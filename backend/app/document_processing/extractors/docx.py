from __future__ import annotations

from io import BytesIO
from typing import Optional

from docx import Document as DocxDocument

from app.config import settings
from app.document.exceptions import MalformedDocument, TextExtractionFailed
from app.document_processing.normalizers.text import count_words, normalize_text
from app.document_processing.schemas import ExtractedPage, ExtractedSection, ExtractionResult


class DocxTextExtractor:
    extractor_name = "python_docx_extractor"
    extractor_version = "1.0"

    def extract(self, content: bytes) -> ExtractionResult:
        try:
            document = DocxDocument(BytesIO(content))
        except Exception as error:
            raise MalformedDocument("DOCX file is malformed.") from error

        raw_blocks = extract_docx_blocks(document)
        raw_text = "\n".join(raw_blocks)
        normalized_text = normalize_text(raw_text)

        if len(normalized_text) > settings.document_max_extracted_characters:
            raise TextExtractionFailed("Extracted text exceeds the configured limit.")

        page = ExtractedPage(
            page_number=1,
            raw_text=raw_text,
            normalized_text=normalized_text,
            character_count=len(normalized_text),
            word_count=count_words(normalized_text),
            has_text=bool(normalized_text),
        )

        return ExtractionResult(
            extraction_method="docx_text",
            extractor_name=self.extractor_name,
            extractor_version=self.extractor_version,
            raw_text=raw_text,
            normalized_text=normalized_text,
            pages=[page],
            sections=build_docx_sections(document),
            metadata={"source": "docx"},
        )


def extract_docx_blocks(document) -> list[str]:
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]

            if cells:
                blocks.append(" | ".join(cells))

    return blocks


def build_docx_sections(document) -> list[ExtractedSection]:
    sections: list[ExtractedSection] = []
    current_heading: Optional[str] = None
    current_blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)

        if not text:
            continue

        style_name = (paragraph.style.name or "").lower() if paragraph.style else ""

        if style_name.startswith("heading"):
            if current_blocks:
                sections.append(
                    make_section(
                        index=len(sections) + 1,
                        heading=current_heading,
                        blocks=current_blocks,
                    )
                )

            current_heading = text
            current_blocks = []
            continue

        current_blocks.append(text)

    if current_blocks:
        sections.append(
            make_section(
                index=len(sections) + 1,
                heading=current_heading,
                blocks=current_blocks,
            )
        )

    return sections


def make_section(
    *,
    index: int,
    heading: Optional[str],
    blocks: list[str],
) -> ExtractedSection:
    return ExtractedSection(
        section_index=index,
        heading=heading,
        section_type="docx_heading" if heading else "body",
        normalized_text=normalize_text("\n".join(blocks)),
        page_start=1,
        page_end=1,
        confidence=80 if heading else 60,
    )
